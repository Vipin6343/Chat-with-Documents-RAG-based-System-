import asyncio
import io
import json
from concurrent.futures import ThreadPoolExecutor

import pymupdf
import pytest
from docx import Document as WordDocument
from fastapi.testclient import TestClient
from sqlalchemy import select
from starlette.datastructures import UploadFile

from app.main import create_app
from app.models import Chunk, Document, Message, VectorIndex
from app.schemas import ChatRequest
from app.services.chat import CitationError, UNSUPPORTED, prepare, stream_answer, validate_citations
from app.services.parsing import parse_chunks
from app.services.storage import ReindexRequired
from conftest import workspace, upload, events


def test_pdf_page_metadata(setup, tmp_path):
    client, store, _, settings = setup
    path = tmp_path / 'pages.pdf'
    with pymupdf.open() as pdf:
        for text in ('apple first page ' * 20, 'banana second page ' * 20):
            page = pdf.new_page()
            page.insert_text((72, 72), text)
        pdf.save(path)
    chunks = parse_chunks(path, 'doc', settings)
    assert {c[2]['page'] for c in chunks} == {1, 2}
    assert all('banana' not in c[1] for c in chunks if c[2]['page'] == 1)
    assert chunks == parse_chunks(path, 'doc', settings)
    assert len({c[0] for c in chunks}) == len(chunks)


def test_docx_and_txt_locations(setup, tmp_path):
    settings = setup[3]
    path = tmp_path / 'notes.docx'
    doc = WordDocument()
    doc.add_paragraph('apple first')
    doc.add_paragraph('banana second')
    doc.save(path)
    assert [c[2] for c in parse_chunks(path, 'd', settings)] == [{'paragraph': 1}, {'paragraph': 2}]
    path = tmp_path / 'notes.txt'
    path.write_text('apple\n\nbanana', encoding='utf-8')
    assert [c[2] for c in parse_chunks(path, 'd', settings)] == [{'paragraph': 1}, {'paragraph': 2}]


def test_isolation_duplicates_and_delete(setup):
    client, store, _, _ = setup
    first, second = workspace(client), workspace(client)
    doc = upload(client, first)['document']
    duplicate = upload(client, first, name='renamed.txt')
    assert duplicate['duplicate'] and duplicate['document']['id'] == doc['id']
    assert not upload(client, second)['duplicate']
    chunks = store.retrieve(first, 'apple')
    assert chunks and all(c['workspace_id'] == first for c in chunks)
    cid = chunks[0]['chunk_id']
    assert client.get(f'/api/workspaces/{second}/sources/{cid}').status_code == 404
    assert client.delete(f'/api/workspaces/{second}/documents/{doc["id"]}').status_code == 404
    conversation = client.post(f'/api/workspaces/{first}/conversations').json()['id']
    assert client.get(f'/api/workspaces/{second}/conversations/{conversation}/messages').status_code == 404
    assert client.post(f'/api/workspaces/{second}/chat/stream', json={'question': 'x', 'conversation_id': conversation}).status_code == 404
    assert client.delete(f'/api/workspaces/{first}/documents/{doc["id"]}').status_code == 204
    assert store.retrieve(first, 'apple') == []
    assert store.retrieve(second, 'apple')
    assert client.get(f'/api/workspaces/{first}/sources/{cid}').status_code == 404
    assert not (store.root / 'uploads' / first / (doc['id'] + '.txt')).exists()


@pytest.mark.parametrize('answer', ['No citation', 'Claim [S99]', 'Claim [S1, S2]', 'Claim [Sx]', 'Claim [S01]'])
def test_invalid_citations(answer):
    with pytest.raises(CitationError):
        validate_citations(answer, [{'source_id': 'S1'}])


def test_citations_and_no_context(setup):
    client, _, provider, _ = setup
    wid = workspace(client)
    result = events(client.post(f'/api/workspaces/{wid}/chat/stream', json={'question': 'unknown'}))
    assert [e[0] for e in result] == ['sources', 'token', 'sources', 'done']
    assert result[1][1]['text'] == UNSUPPORTED
    assert not provider.prompts
    assert validate_citations('Claim [S1]', [{'source_id': 'S1'}, {'source_id': 'S2'}]) == [{'source_id': 'S1'}]


def test_stream_success_history_and_source_metadata(setup):
    client, _, provider, _ = setup
    wid = workspace(client)
    upload(client, wid, name='actual.txt')
    response = client.post(f'/api/workspaces/{wid}/chat/stream', json={'question': 'apple?'})
    result = events(response)
    assert result[-1][0] == 'done'
    final_sources = result[-2][1]['sources']
    assert final_sources[0]['filename'] == 'actual.txt'
    cid, mid = result[-1][1]['conversation_id'], result[-1][1]['message_id']
    messages = client.get(f'/api/workspaces/{wid}/conversations/{cid}/messages').json()
    answer = next(m for m in messages if m['id'] == mid)
    assert answer['status'] == 'complete' and answer['sources'] == final_sources
    assert answer['content'] == provider.answer
    source = client.get(f'/api/workspaces/{wid}/sources/{final_sources[0]["chunk_id"]}').json()
    assert source['text'] == 'apple is red'
    followup = events(client.post(f'/api/workspaces/{wid}/chat/stream', json={'question': 'What color is it?', 'conversation_id': cid}))
    assert followup[-1][0] == 'done' and provider.rewrites


@pytest.mark.parametrize('failure', ['provider', 'citation'])
def test_stream_failure_never_done(setup, failure):
    client, _, provider, _ = setup
    wid = workspace(client)
    upload(client, wid)
    provider.fail = failure == 'provider'
    if failure == 'citation':
        provider.answer = 'Invented claim [S99]'
    result = events(client.post(f'/api/workspaces/{wid}/chat/stream', json={'question': 'apple?'}))
    assert result[-1][0] == 'error' and 'done' not in [r[0] for r in result]
    assert 'secret' not in json.dumps(result)
    ids = result[-1][1]
    messages = client.get(f'/api/workspaces/{wid}/conversations/{ids["conversation_id"]}/messages').json()
    assert next(m for m in messages if m['id'] == ids['message_id'])['status'] == 'failed'


def test_upload_validation_and_failed_state(setup):
    client, store, _, settings = setup
    wid = workspace(client)
    assert upload(client, wid, name='evil.exe')['error']
    assert upload(client, wid, text=b'')['error']
    for name, data in [('bad.pdf', b'not pdf'), ('bad.txt', b'\xff'), ('empty.txt', b'   ')]:
        result = upload(client, wid, text=data, name=name)
        with store.session() as db:
            doc = db.get(Document, result['document']['id'])
            assert doc.status == 'failed' and doc.error
    settings.max_file_bytes = 3
    assert upload(client, wid, text=b'too large')['error']
    settings.max_files = 1
    assert client.post(f'/api/workspaces/{wid}/documents', files=[('files', ('a.txt', b'a')), ('files', ('b.txt', b'b'))]).status_code == 413


def test_reload_and_incompatible_settings(setup):
    client, store, providers, settings = setup
    wid = workspace(client)
    upload(client, wid)
    from app.services.storage import Store
    second = Store(settings, providers)
    try:
        assert second.retrieve(wid, 'apple')
        settings.embedding_model = 'changed-model'
        with pytest.raises(ReindexRequired):
            second.retrieve(wid, 'apple')
        second.reindex(wid)
        assert second.retrieve(wid, 'apple')
    finally:
        second.engine.dispose()


def test_concurrent_ingestion_keeps_both_documents(setup):
    client, store, _, _ = setup
    wid = workspace(client)
    def ingest(text):
        _, pending = store.upload(wid, [UploadFile(io.BytesIO(text), filename='facts.txt')])
        store.process(wid, pending[0])
    with ThreadPoolExecutor(max_workers=2) as executor:
        list(executor.map(ingest, [b'apple red', b'banana yellow']))
    assert len({c['document_id'] for c in store.retrieve(wid, 'apple banana')}) == 2


def test_restart_marks_unfinished_work(setup):
    client, store, _, _ = setup
    wid = workspace(client)
    _, pending = store.upload(wid, [UploadFile(io.BytesIO(b'apple'), filename='a.txt')])
    cid, mid, _ = prepare(store, wid, ChatRequest(question='apple?'))
    store.recover()
    with store.session() as db:
        assert db.get(Document, pending[0]).status == 'failed'
        assert db.get(Message, mid).status == 'interrupted'


def test_disconnect_persists_interrupted(setup):
    client, store, _, _ = setup
    wid = workspace(client)
    cid, mid, history = prepare(store, wid, ChatRequest(question='apple?'))
    class Disconnected:
        async def is_disconnected(self):
            return True
    async def consume():
        return [event async for event in stream_answer(store, Disconnected(), wid, 'apple?', cid, mid, history)]
    assert asyncio.run(consume()) == []
    with store.session() as db:
        assert db.get(Message, mid).status == 'interrupted'


def test_second_server_rejected(setup):
    _, _, providers, settings = setup
    import portalocker
    with pytest.raises(portalocker.exceptions.LockException):
        with TestClient(create_app(settings, providers)):
            pass


def test_deleted_source_during_stream_fails(setup):
    client, store, provider, _ = setup
    wid = workspace(client)
    did = upload(client, wid)['document']['id']
    original = provider.astream
    async def deleting(prompt):
        async for token in original(prompt):
            yield token
        store.delete_document(wid, did)
    provider.astream = deleting
    result = events(client.post(f'/api/workspaces/{wid}/chat/stream', json={'question': 'apple?'}))
    assert result[-1][0] == 'error'
    assert not any(e[0] == 'done' for e in result)


def test_dimension_mismatch_and_corrupt_index(setup):
    client, store, provider, _ = setup
    wid = workspace(client)
    upload(client, wid)
    original = provider.embeddings
    class ChangedEmbeddings:
        def embed_query(self, text):
            return [1., 2.]
    provider.embeddings = lambda: ChangedEmbeddings()
    with pytest.raises(ReindexRequired, match='dimensions'):
        store.retrieve(wid, 'apple')
    provider.embeddings = original
    with store.session.begin() as db:
        db.get(VectorIndex, wid).payload = b'corrupt'
    with pytest.raises(ReindexRequired, match='checksum'):
        store.retrieve(wid, 'apple')


def test_provider_errors_are_not_persisted_verbatim(setup):
    client, store, provider, _ = setup
    wid = workspace(client)
    class BrokenEmbeddings:
        def embed_documents(self, texts):
            raise ValueError('secret api key and document contents')
    provider.embeddings = lambda: BrokenEmbeddings()
    did = upload(client, wid)['document']['id']
    with store.session() as db:
        doc = db.get(Document, did)
        assert doc.status == 'failed'
        assert 'secret' not in doc.error
        assert not list(db.scalars(select(Chunk).where(Chunk.document_id == did)))


def test_context_budget_and_conversation_busy(setup):
    client, store, _, settings = setup
    wid = workspace(client)
    settings.context_window = 2048
    response = events(client.post(f'/api/workspaces/{wid}/chat/stream', json={'question': 'x' * 2000}))
    assert response[-1][0] == 'error'
    cid, _, _ = prepare(store, wid, ChatRequest(question='apple?'))
    assert client.post(f'/api/workspaces/{wid}/chat/stream', json={'question': 'banana?', 'conversation_id': cid}).status_code == 409


def test_scanned_pdf_and_safe_filename(setup, tmp_path):
    client, store, _, _ = setup
    wid = workspace(client)
    path = tmp_path / 'blank.pdf'
    with pymupdf.open() as pdf:
        pdf.new_page()
        pdf.save(path)
    did = upload(client, wid, text=path.read_bytes(), name='blank.pdf')['document']['id']
    with store.session() as db:
        assert 'OCR' in db.get(Document, did).error
    result = upload(client, wid, name='../../outside.txt')
    with store.session() as db:
        stored = store.path(db.get(Document, result['document']['id']))
        assert stored.parent == store.root / 'uploads' / wid
        assert stored.name == result['document']['id'] + '.txt'


def test_body_limit_including_chunked_requests(setup):
    _, _, provider, settings = setup
    from app.main import BodyLimitMiddleware
    from fastapi import FastAPI, Request
    small = FastAPI()
    small.add_middleware(BodyLimitMiddleware, limit=10)
    @small.post('/')
    async def body(request: Request):
        return {'size': len(await request.body())}
    with TestClient(small) as client:
        response = client.post('/', content=iter([b'123456', b'123456']))
        assert response.status_code == 413


def test_stream_cancellation_after_token(setup):
    client, store, provider, _ = setup
    wid = workspace(client)
    upload(client, wid)
    cid, mid, history = prepare(store, wid, ChatRequest(question='apple?'))
    from langchain_core.messages import AIMessageChunk
    async def cancelled(prompt):
        yield AIMessageChunk(content='partial')
        raise asyncio.CancelledError()
    provider.astream = cancelled
    class Connected:
        async def is_disconnected(self):
            return False
    async def consume():
        collected = []
        with pytest.raises(asyncio.CancelledError):
            async for event in stream_answer(store, Connected(), wid, 'apple?', cid, mid, history):
                collected.append(event)
        assert any('event: token' in e for e in collected)
        assert not any('event: done' in e for e in collected)
    asyncio.run(consume())
    with store.session() as db:
        message = db.get(Message, mid)
        assert message.status == 'interrupted' and message.content == 'partial'


def test_real_provider_constructors_without_network():
    from app.config import Settings
    from app.services.providers import Providers
    for provider, model in [('ollama', 'llama3.2'), ('openai', 'gpt-4.1-mini')]:
        settings = Settings(_env_file=None, llm_provider=provider, embedding_provider=provider,
                            llm_model=model, openai_api_key='test-not-a-real-key')
        factory = Providers(settings)
        assert factory.llm() is not None
        assert factory.embeddings() is not None
