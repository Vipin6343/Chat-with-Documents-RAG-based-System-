import re
import zipfile
from pathlib import Path
from uuid import NAMESPACE_URL, uuid5

import pymupdf
from docx import Document
from langchain_core.documents import Document as TextDocument
from langchain_text_splitters import RecursiveCharacterTextSplitter


def parse_chunks(path: Path, document_id: str, settings):
    sections = []
    total = 0

    def add(text, location):
        nonlocal total
        total += len(text)
        if total > settings.max_extracted_chars:
            raise ValueError('Extracted text exceeds configured character limit')
        if text.strip():
            sections.append(TextDocument(page_content=text, metadata=location))

    if path.suffix == '.pdf':
        with pymupdf.open(path) as pdf:
            if pdf.needs_pass:
                raise ValueError('Password-protected PDFs are unsupported')
            for number, page in enumerate(pdf, 1):
                add(page.get_text(), {'page': number})
    elif path.suffix == '.docx':
        with zipfile.ZipFile(path) as archive:
            if sum(i.file_size for i in archive.infolist()) > 50_000_000:
                raise ValueError('Expanded DOCX exceeds 50 MB limit')
        doc = Document(path)
        for number, paragraph in enumerate(doc.paragraphs, 1):
            add(paragraph.text, {'paragraph': number})
        for number, table in enumerate(doc.tables, 1):
            for row_number, row in enumerate(table.rows, 1):
                add(' | '.join(c.text for c in row.cells), {'table': number, 'row': row_number})
    else:
        for number, paragraph in enumerate(re.split(r'\n\s*\n', path.read_text(encoding='utf-8-sig')), 1):
            add(paragraph, {'paragraph': number})
    if not sections:
        raise ValueError('No extractable text. Scanned PDFs require OCR, outside this version.')
    splitter = RecursiveCharacterTextSplitter(chunk_size=settings.chunk_size, chunk_overlap=settings.chunk_overlap)
    return [(str(uuid5(NAMESPACE_URL, f'{document_id}:{i}')), d.page_content, d.metadata)
            for i, d in enumerate(splitter.split_documents(sections))]
